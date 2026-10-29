"""Filling Word templates safely, checking required fields, padding logos, signatures and converting to PDF."""
import html, io, os, pathlib, shutil, subprocess, tempfile, zipfile, re
from docxtpl import DocxTemplate, InlineImage
from docx.shared import Mm
from jinja2.sandbox import SandboxedEnvironment
from PIL import Image
from .fields import REQUIRED

LOGO_RATIO = 1.6                     # every template's logo frame is about 1.6 wide : 1 high
BANNER = "NOT AUTHORISED FOR USE UNTIL SIGNED"
SIGNATURE_HEIGHT_MM = 14

def sandbox() -> SandboxedEnvironment:
    # autoescape keeps "Smith & Sons" valid XML; None renders as blank
    return SandboxedEnvironment(autoescape=True, finalize=lambda v: "" if v is None else v)

def pad_logo(image_bytes: bytes, ratio: float = LOGO_RATIO, max_w: int = 800) -> bytes:
    """Any uploaded image -> PNG padded (never stretched) to the template frame's proportions."""
    im = Image.open(io.BytesIO(image_bytes)).convert("RGBA")
    im.thumbnail((max_w, max_w))
    w, h = im.size
    if w / h > ratio: cw, ch = w, round(w / ratio)
    else: cw, ch = round(h * ratio), h
    canvas = Image.new("RGBA", (cw, ch), (255, 255, 255, 0))
    canvas.paste(im, ((cw - w) // 2, (ch - h) // 2), im)
    out = io.BytesIO(); canvas.save(out, "PNG"); return out.getvalue()

def blank_logo() -> bytes:
    out = io.BytesIO(); Image.new("RGBA", (16, 10), (255, 255, 255, 0)).save(out, "PNG"); return out.getvalue()

def clean_signature(image_bytes: bytes) -> bytes:
    """A drawn or photographed signature -> a trimmed PNG with a see-through background, so it sits on any form."""
    try:
        im = Image.open(io.BytesIO(image_bytes)); im.load()
    except Exception:
        raise ValueError("That file isn't a picture. Draw your signature, or upload a photo of it.")
    im = im.convert("RGBA"); im.thumbnail((1800, 1800))
    im.putdata([(r, g, b, 0) if r > 200 and g > 200 and b > 200 else (r, g, b, a) for (r, g, b, a) in im.getdata()])
    box = im.getchannel("A").getbbox()
    if not box:
        raise ValueError("The signature is blank. Draw it or upload it again.")
    im = im.crop(box); im.thumbnail((900, 300))
    out = io.BytesIO(); im.save(out, "PNG"); return out.getvalue()

def has_logo_picture(docx_bytes: bytes) -> bool:
    z = zipfile.ZipFile(io.BytesIO(docx_bytes))
    return any(re.search(rb'<pic:cNvPr[^>]*\bname="logo"', z.read(n)) for n in z.namelist()
               if n.startswith("word/") and n.endswith(".xml"))

def render(template_bytes: bytes, context: dict, logo_png: bytes | None = None, images: dict | None = None) -> bytes:
    """Fills a template. `images` puts a picture (a saved signature) where that field would otherwise print text."""
    tpl = DocxTemplate(io.BytesIO(template_bytes))
    if has_logo_picture(template_bytes):
        tpl.replace_pic("logo", io.BytesIO(logo_png or blank_logo()))
    ctx = dict(context)
    for k, png in (images or {}).items():
        if png:
            ctx[k] = InlineImage(tpl, io.BytesIO(png), height=Mm(SIGNATURE_HEIGHT_MM))
    tpl.render(ctx, jinja_env=sandbox())
    out = io.BytesIO(); tpl.save(out); return out.getvalue()

def add_signature_block(docx_bytes: bytes, png: bytes | None, lines: list[str]) -> bytes:
    """A dated "Signed electronically" block at the end of a document. Every signed RAMS part gets one; a part edited
    by hand (whose signature box is gone) also gets the signature image in it."""
    from docx import Document
    from docx.shared import Pt
    doc = Document(io.BytesIO(docx_bytes))
    head = doc.add_paragraph(); r = head.add_run("Signed electronically"); r.bold = True
    if png:
        doc.add_paragraph().add_run().add_picture(io.BytesIO(png), height=Mm(SIGNATURE_HEIGHT_MM))
    for line in lines:
        run = doc.add_paragraph().add_run(line); run.font.size = Pt(9)
    out = io.BytesIO(); doc.save(out); return out.getvalue()

def add_banner(docx_bytes: bytes) -> bytes:
    """The unsigned-handover banner is a safety control: forced into uploaded handover templates."""
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    doc = Document(io.BytesIO(docx_bytes))
    first = doc.paragraphs[0] if doc.paragraphs else doc.add_paragraph()
    p = first.insert_paragraph_before()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(BANNER); r.bold = True; r.font.size = Pt(16); r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    out = io.BytesIO(); doc.save(out); return out.getvalue()

def text_of(docx_bytes: bytes) -> str:
    z = zipfile.ZipFile(io.BytesIO(docx_bytes))
    parts = [n for n in z.namelist() if n.startswith("word/") and n.endswith(".xml") and "_rels" not in n]
    return html.unescape("\n".join(re.sub(r"<[^>]+>", "", z.read(n).decode("utf8", "ignore")) for n in parts))

def missing_required(kind: str, context: dict, *, logo_present: bool = True) -> list[str]:
    missing = [f for f in REQUIRED.get(kind, []) if context.get(f) in (None, "", [])]
    if kind == "inspection":
        for i, c in enumerate(context.get("checklist") or [], 1):
            if not c.get("result"):
                missing.append(f"checklist item {i} ({c.get('item')}) has no result")
    if not logo_present:
        missing.append("firm logo (upload it in Settings)")
    return missing

class PdfConversionError(RuntimeError):
    """LibreOffice is installed but this conversion failed or timed out."""

# Where the Windows installer puts LibreOffice. It does not add itself to the system path, so look here too.
WINDOWS_DEFAULTS = [r"C:\Program Files\LibreOffice\program\soffice.exe",
                    r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"]

def soffice_path() -> str | None:
    """SOFFICE_PATH if set and real, then the system path (Railway's Docker image), then the Windows install folder."""
    env = os.environ.get("SOFFICE_PATH")
    if env and os.path.exists(env):
        return env
    found = shutil.which("soffice") or shutil.which("libreoffice")
    if found:
        return found
    if os.name == "nt":
        for p in WINDOWS_DEFAULTS:
            if os.path.exists(p):
                return p
    return None

def to_pdf(docx_bytes: bytes, timeout: int = 120) -> bytes | None:
    exe = soffice_path()
    if not exe:
        return None                                    # caller reports "PDF unavailable on this machine"
    with tempfile.TemporaryDirectory() as d:
        src = os.path.join(d, "doc.docx")
        with open(src, "wb") as fh: fh.write(docx_bytes)
        # A proper file URI on every OS: file:///C:/Users/... on Windows, file:///tmp/... on Linux.
        # A malformed one makes LibreOffice report "bootstrap.ini is corrupt" and exit.
        profile = pathlib.Path(d, "profile").as_uri()
        cmd = [exe, f"-env:UserInstallation={profile}", "--headless", "--norestore", "--nolockcheck",
               "--convert-to", "pdf", "--outdir", d, src]
        try:
            r = subprocess.run(cmd, timeout=timeout, capture_output=True)
        except subprocess.TimeoutExpired:
            raise PdfConversionError(f"PDF conversion timed out after {timeout} seconds")
        pdf_path = os.path.join(d, "doc.pdf")
        if r.returncode != 0 or not os.path.exists(pdf_path):
            detail = (r.stderr or r.stdout or b"").decode("utf8", "ignore").strip()[:500]
            raise PdfConversionError(f"PDF conversion failed (LibreOffice exit code {r.returncode}): {detail or 'no output'}")
        with open(pdf_path, "rb") as fh: return fh.read()
