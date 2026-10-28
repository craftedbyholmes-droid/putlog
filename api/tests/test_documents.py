"""Acceptance tests 6, 15, 16, 17, 18, 20 (synthetic), 21 and 22 at the document level."""
import io, re, zipfile, hashlib, pathlib, pytest
from docx import Document
from app import documents as D, template_checks as T, quotes
from app.fields import INSPECTION_CHECKLIST
from conftest import FIRM, JOB, ROOT

def leftover(docx): return re.findall(r"\{\{|\{%|\}\}|%\}", D.text_of(docx))

def logo_replaced(docx, logo_png):
    z = zipfile.ZipFile(io.BytesIO(docx)); want = hashlib.md5(logo_png).hexdigest()
    hits = []
    for n in z.namelist():
        if n.startswith("word/") and n.endswith(".xml") and "_rels" not in n:
            x = z.read(n).decode()
            for rid in re.findall(r'<pic:cNvPr[^>]*name="logo".*?r:embed="(rId[^"]+)"', x, re.S):
                rels = z.read(n.replace("word/", "word/_rels/") + ".rels").decode()
                tgt = re.search(rf'Id="{rid}"[^>]*Target="([^"]+)"', rels) or re.search(rf'Target="([^"]+)"[^>]*Id="{rid}"', rels)
                hits.append(hashlib.md5(z.read("word/" + tgt.group(1))).hexdigest() == want)
    return hits and all(hits)

def test_logo_padded_never_stretched(square_logo):
    from PIL import Image
    png = D.pad_logo(square_logo); im = Image.open(io.BytesIO(png))
    assert im.format == "PNG" and abs(im.width / im.height - 1.6) < 0.01

@pytest.mark.parametrize("kind", ["Method_Statement", "Risk_Assessment"])
def test_15_rams_from_standard_template(std, square_logo, kind):
    logo = D.pad_logo(square_logo)
    out = D.render(std(kind), {**FIRM, **JOB}, logo)
    assert leftover(out) == []
    text = D.text_of(out)
    assert "Acme & Co Builders" in text and "14 Church Street" in text and "NASC TG20:21" in text
    assert logo_replaced(out, logo)
    assert D.missing_required(kind.lower(), {**FIRM, **JOB}) == []
    assert "supervisor_name" in D.missing_required("method_statement", {**FIRM, **JOB, "supervisor_name": ""})
    pdf = D.to_pdf(out)
    assert pdf is None or pdf[:4] == b"%PDF"

def quote_context(vat):
    lines = quotes.lines_for_scaffold({"name": "Front", "pricing_mode": "squareage"},
                                      [{"quoted_area_m2": 60}, {"quoted_area_m2": 48}], 3.5)
    t = quotes.totals(lines, vat)
    g = quotes.gbp
    return {**FIRM, "client_name": "Acme & Co Builders", "site_address": "14 Church Street, Cromer", "quote_number": "Q0001",
            "quote_date": "29/10/2026", "valid_until": "28/11/2026", "valid_days": 30, "job_reference": "JOB0001",
            "hire_included_weeks": 4, "prepared_by_name": "Keith",
            "elevations": [{"scaffold": "Front", "label": "A", "length": "10.00", "height": "6.00", "area": "60.00"},
                           {"scaffold": "Front", "label": "B", "length": "8.00", "height": "6.00", "area": "48.00"}],
            "lines": [{**l, "unit_price": g(l["unit_price"]), "line_total": g(l["line_total"])} for l in lines],
            "subtotal": g(t["subtotal"]), "vat_registered": vat, "vat_number": "GB123456789" if vat else "",
            "vat_rate": t["vat_rate"], "vat_amount": g(t["vat_amount"]), "total": g(t["total"])}

@pytest.mark.parametrize("vat", [False, True])
def test_06_quote_document(std, square_logo, vat):
    out = D.render(std("Quote"), quote_context(vat), D.pad_logo(square_logo))
    assert leftover(out) == []
    text = D.text_of(out)
    assert "Smith & Sons Scaffolding Ltd" in text and "Q0001" in text and "60.00" in text and "48.00" in text
    assert ("£453.60" if vat else "£378.00") in text
    assert ("GB123456789" in text) == vat and (("VAT at" in text) == vat)

def handover_ctx(pending):
    return {**FIRM, **JOB, "certificate_ref": "HC0001", "scaffold_type": "Tied independent tube and fitting",
            "permitted_use": "Access and light duty", "loading_class": "Class 2 (1.5 kN/m²)",
            "max_loaded_platforms": "1 fully loaded + 1 at 50%", "restrictions": "None", "max_leg_load_kn": "12.5",
            "add_ons_permitted": "None", "inspecting_party": "Us (the scaffold contractor)",
            "inspection_responsible_name": "Dan Price", "inspection_responsible_phone": "07700 900001",
            "next_inspection_due": "09/11/2026", "handover_date": "02/11/2026", "issued_by_name": "Keith",
            "issued_by_position": "Owner", "issued_by_signature": "", "issued_at": "02/11/2026 16:10",
            "confirmed_by_name": "" if pending else "Jo Bloggs", "confirmed_by_role": "" if pending else "Site manager",
            "confirmed_by_company": "" if pending else "Acme", "confirmed_by_signature": "", "confirmed_at": "",
            "acceptance_method_text": "", "signed_ip": "", "signature_pending": pending}

def test_21_handover_banner_and_required(std, square_logo):
    logo = D.pad_logo(square_logo)
    pending = D.render(std("Handover_Certificate"), handover_ctx(True), logo)
    signed = D.render(std("Handover_Certificate"), handover_ctx(False), logo)
    assert D.BANNER in D.text_of(pending).upper() and D.BANNER not in D.text_of(signed).upper()
    assert leftover(pending) == [] and logo_replaced(pending, logo)
    assert D.missing_required("handover", handover_ctx(True)) == []
    assert D.missing_required("handover", {**handover_ctx(True), "loading_class": ""}) == ["loading_class"]

def test_21_banner_forced_into_uploaded_handover():
    d = Document(); d.add_paragraph("Our own handover certificate for {{ client_name }}")
    b = io.BytesIO(); d.save(b)
    out = D.add_banner(D.render(b.getvalue(), {"client_name": "Acme"}))
    assert D.text_of(out).index(D.BANNER) < D.text_of(out).index("Our own handover")

def inspection_ctx(defects):
    return {**FIRM, **JOB, "report_ref": "IR0001", "inspection_type": "7-day", "inspected_at": "09/11/2026 08:30",
            "previous_inspection_at": "02/11/2026", "weather": "Dry, light wind", "scaffold_ref": "Front elevation",
            "scaffold_type": "Tied independent", "loading_class": "Class 2", "outcome": "Safe to use",
            "tag_updated": "Yes", "next_inspection_due": "16/11/2026", "copies_sent_to": "Office; client",
            "signed_at": "", "inspector_name": "Dan Price", "inspector_position": "Advanced scaffolder",
            "inspector_qualification": "CISRS Advanced", "inspector_signature": "",
            "checklist": [{"item": i, "result": "Yes", "comment": ""} for i in INSPECTION_CHECKLIST],
            "defects": defects}

def test_22_inspection_report(std):
    none = D.render(std("Inspection_Report"), inspection_ctx([]))
    some = D.render(std("Inspection_Report"), inspection_ctx([{"matter": "Missing toeboard lift 2", "action_taken": "Fitted",
                    "further_action": "None", "responsible": "Dan Price", "due_date": "09/11/2026"}]))
    for out in (none, some):
        assert leftover(out) == []
        text = D.text_of(out)
        for item in INSPECTION_CHECKLIST: assert item in text
        for f in ("Acme & Co Builders", "2 Market Place", "14 Church Street", "Front elevation", "09/11/2026 08:30", "Dan Price", "Advanced scaffolder"):
            assert f in text
    assert "no matter" in D.text_of(none).lower() and "Missing toeboard" in D.text_of(some)
    assert "no matter" not in D.text_of(some).lower()
    bad = inspection_ctx([]); bad["checklist"][3]["result"] = ""
    assert any("Bracing" in m for m in D.missing_required("inspection", bad))

# ---------- uploads ----------
def docx_with(*paras, header=None):
    d = Document()
    for p in paras: d.add_paragraph(p)
    if header: d.sections[0].header.paragraphs[0].text = header
    b = io.BytesIO(); d.save(b); return b.getvalue()

def test_16_valid_upload_is_accepted():
    r = T.check_upload("ours.docx", docx_with("Client: {{ client_name }}", "Site: {{ site_address }}", "Date: {{ date }}",
                                              "Area: {{ work_area }}", "By: {{ prepared_by_name }}"))
    assert r.ok and r.tags == ["client_name", "date", "prepared_by_name", "site_address", "work_area"]
    assert any("logo" in w for w in r.warnings)                     # no picture named logo: warning only

def test_17_bad_uploads_named_problems():
    assert "not a .docx" in T.check_upload("x.pdf", b"%PDF").errors[0]
    assert "limit" in T.check_upload("x.docx", b"0" * (T.MAX_BYTES + 1)).errors[0]
    assert "damaged" in T.check_upload("x.docx", b"not a zip").errors[0]
    # macros
    src = docx_with("{{ client_name }}"); buf = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(src)) as zin, zipfile.ZipFile(buf, "w") as zout:
        for n in zin.namelist(): zout.writestr(n, zin.read(n))
        zout.writestr("word/vbaProject.bin", b"x")
    assert any("macros" in e for e in T.check_upload("m.docx", buf.getvalue()).errors)
    # external link
    d = Document(); p = d.add_paragraph("{{ client_name }}")
    d.part.relate_to("https://evil.example/x", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    b = io.BytesIO(); d.save(b)
    assert any("evil.example" in e for e in T.check_upload("l.docx", b.getvalue()).errors)
    # unknown tag
    assert any("Unknown field {{ customer }}" in e for e in T.check_upload("u.docx", docx_with("{{ customer }}")).errors)
    # missing core fields: warning, still ok
    r = T.check_upload("c.docx", docx_with("{{ client_name }}"))
    assert r.ok and any("Missing core fields" in w for w in r.warnings)

@pytest.mark.parametrize("evil", ["{{ client_name.__class__ }}", "{{ 7*7 }}", "{% for x in range(9) %}x{% endfor %}",
                                  "{{ config }}", "{% include '/etc/passwd' %}", "{{ cycler.__init__.__globals__.os }}",
                                  "{{ lipsum('http://x') }}", "{# comment #}", "{{ client_name|upper }}"])
def test_18_sandbox_rejects_anything_but_plain_fields(evil):
    r = T.check_upload("e.docx", docx_with("Client: {{ client_name }}", evil))
    assert not r.ok

def test_18_render_runs_in_sandbox():
    from jinja2.exceptions import SecurityError
    with pytest.raises(SecurityError):
        D.render(docx_with("{{ x.__class__.__mro__ }}"), {"x": "a"})

def test_20_untagged_labels_proposed_and_inserted():
    d = Document()
    d.sections[0].header.paragraphs[0].text = "Customer :            Date :"
    d.add_paragraph("Site :")
    d.add_paragraph("Work Area :          Quotation No :")
    d.add_paragraph("Grid Ref :")
    t = d.add_table(rows=1, cols=2); t.rows[0].cells[0].text = "Supervisor :"
    d.add_paragraph("Item No : ")
    b = io.BytesIO(); d.save(b); src = b.getvalue()
    proposals, unmatched = T.find_labels(src)
    assert {p["field"] for p in proposals} == {"client_name", "date", "site_address", "work_area", "quote_reference",
                                               "supervisor_name", "item_no"}
    assert [u["label"] for u in unmatched] == ["Grid Ref"]
    tagged = T.insert_tags(src, proposals)
    assert T.check_upload("t.docx", tagged).ok
    out = D.render(tagged, {**FIRM, **JOB})
    hdr = Document(io.BytesIO(out)).sections[0].header.paragraphs[0].text
    assert "Acme & Co Builders" in hdr and "29/10/2026" in hdr
    body = D.text_of(out)
    assert "Dan Price" in body and "Q0001" in body and "Front elevation" in body

ORIGINALS = sorted((ROOT / "fixtures" / "private").glob("*.docx"))
@pytest.mark.skipif(not ORIGINALS, reason="put the two original RAMS files in fixtures/private to run test 20 on them")
@pytest.mark.parametrize("path", ORIGINALS, ids=lambda p: p.name)
def test_20_original_documents(path):
    proposals, unmatched = T.find_labels(path.read_bytes())
    found = {p["field"] for p in proposals}
    assert {"client_name", "site_address", "date"} <= found
    out = D.render(T.insert_tags(path.read_bytes(), proposals), {**FIRM, **JOB})
    assert "Acme & Co Builders" in D.text_of(out)
